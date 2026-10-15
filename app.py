import re
import textwrap
import calendar
import json
import os
from datetime import datetime, timedelta, timezone
import traceback
import pandas as pd
import streamlit as st
from io import BytesIO, StringIO

st.set_page_config(page_title="Distribution", layout="wide")

T = {
    "title": "\u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u062a\u0648\u0632\u064a\u0639 \u062d\u0633\u0628 \u0627\u0644\u0645\u0643\u062a\u0628",
    "up_agents": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0648\u0643\u0644\u0627\u0621",
    "up_dist": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u062a\u0648\u0632\u064a\u0639",
    "up_quality": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0643\u0648\u0627\u0644\u062a\u064a",
    "sheet": "\u0634\u064a\u062a",
    "agents": "\u0627\u0644\u0648\u0643\u0644\u0627\u0621",
    "dist": "\u0627\u0644\u062a\u0648\u0632\u064a\u0639",
    "quality": "\u0627\u0644\u0643\u0648\u0627\u0644\u062a\u064a",
    "calc": "\u0627\u062d\u0633\u0628",
    "office_name": "\u0627\u0633\u0645 \u0627\u0644\u0645\u0643\u062a\u0628",
    "office_no": "\u0631\u0642\u0645 \u0627\u0644\u0645\u0643\u062a\u0628",
    "item": "\u0627\u0644\u0635\u0646\u0641",
    "total": "\u0627\u0644\u0645\u062c\u0645\u0648\u0639",
    "unmatched": "\u063a\u064a\u0631 \u0645\u0637\u0627\u0628\u0642",
    "empty": "\u0645\u0627 \u0628\u0642\u0649 \u0623\u064a \u0635\u0641 \u0628\u0639\u062f \u0627\u0644\u0641\u0644\u062a\u0631\u0629. \u062a\u0623\u0643\u062f \u0645\u0646 \u0639\u0645\u0648\u062f E \u0648\u0627\u0644\u0623\u0635\u0646\u0627\u0641.",
    "done": "\u062a\u0645. \u0639\u062f\u062f \u0627\u0644\u0645\u0643\u0627\u062a\u0628",
    "unm_rows": "\u0635\u0641\u0648\u0641 \u063a\u064a\u0631 \u0645\u0637\u0627\u0628\u0642\u0629",
    "orig": "\u0635\u0641\u0648\u0641 \u0627\u0644\u062a\u0648\u0632\u064a\u0639 \u0627\u0644\u0623\u0635\u0644\u064a\u0629",
    "dedup": "\u0628\u0639\u062f \u0625\u0632\u0627\u0644\u0629 \u0627\u0644\u0645\u0643\u0631\u0631 (\u0622\u062e\u0631 \u062a\u0627\u0631\u064a\u062e)",
    "real": "\u0627\u0644\u0645\u0648\u0632\u0651\u0639 \u0641\u0639\u0644\u064a\u0627\u064b",
    "bad_dates": "\u062a\u0648\u0627\u0631\u064a\u062e \u063a\u064a\u0631 \u0645\u0642\u0631\u0648\u0621\u0629",
    "download": "\u062a\u062d\u0645\u064a\u0644 \u0627\u0644\u0646\u062a\u064a\u062c\u0629 (\u0625\u0643\u0633\u0644)",
    "file_name": "distribution_result.xlsx",
    "act_red": "\u062a\u0641\u0639\u064a\u0644 \u0643\u0644\u064a \u0631\u064a\u062f",
    "act_use": "\u062a\u0641\u0639\u064a\u0644 \u0643\u0644\u064a \u064a\u0648\u0632",
    "act_net": "\u062a\u0641\u0639\u064a\u0644 \u0643\u0644\u064a \u0646\u062a",
    "last_ru": "\u062a\u0641\u0639\u064a\u0644 \u0627\u0644\u062e\u0637 \u0627\u0644\u064a\u0648\u0645\u064a (\u064a\u0648\u0632 \u0648\u0631\u064a\u062f)",
    "last_net": "\u062a\u0641\u0639\u064a\u0644 \u0627\u0644\u0646\u062a \u0627\u0644\u064a\u0648\u0645\u064a",
    "q_info": "\u0627\u0644\u0643\u0648\u0627\u0644\u062a\u064a: \u0639\u062f\u062f \u0627\u0644\u0635\u0641\u0648\u0641",
    "q_last": "\u0622\u062e\u0631 \u064a\u0648\u0645",
    "q_ignored": "\u0635\u0641\u0648\u0641 \u0645\u062a\u062c\u0627\u0647\u0644\u0629 (\u0646\u0648\u0639 \u063a\u064a\u0631 \u0645\u0639\u0631\u0648\u0641)",
    "q_baddate": "\u062a\u0648\u0627\u0631\u064a\u062e \u0643\u0648\u0627\u0644\u062a\u064a \u063a\u064a\u0631 \u0645\u0642\u0631\u0648\u0621\u0629",
    "pts_word": "\u0646\u0642\u0627\u0637",
    "pts_ru": "\u0646\u0642\u0627\u0637 \u0627\u0644\u0627\u062a\u0635\u0627\u0644 \u0627\u0644\u0643\u0644\u064a\u0629",
    "pts_net": "\u0646\u0642\u0627\u0637 \u0627\u0644\u0646\u062a \u0627\u0644\u0643\u0644\u064a\u0629",
    "pts_last_ru": "\u0646\u0642\u0627\u0637 \u0627\u0644\u0627\u062a\u0635\u0627\u0644 \u0627\u0644\u064a\u0648\u0645\u064a\u0629",
    "pts_last_net": "\u0646\u0642\u0627\u0637 \u0627\u0644\u0646\u062a \u0627\u0644\u064a\u0648\u0645\u064a\u0629",
    "up_inc": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u062d\u0627\u0641\u0632",
    "inc": "\u062d\u0627\u0641\u0632",
    "inc_info": "\u0627\u0644\u062d\u0627\u0641\u0632: \u0639\u062f\u062f \u0627\u0644\u0635\u0641\u0648\u0641",
    "inc_after": "\u0628\u0639\u062f \u0625\u0632\u0627\u0644\u0629 \u0627\u0644\u0635\u0641\u0631",
    "inc_offices": "\u0645\u0643\u0627\u062a\u0628 \u0628\u0627\u0644\u062d\u0627\u0641\u0632",
    "up_wal": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0645\u062d\u0627\u0641\u0638",
    "wal": "\u0645\u062d\u0641\u0638\u0629",
    "wal_info": "\u0627\u0644\u0645\u062d\u0627\u0641\u0638: \u0639\u062f\u062f \u0627\u0644\u0635\u0641\u0648\u0641",
    "wal_after": "\u0628\u0639\u062f \u0641\u0644\u062a\u0631 Y",
    "wal_offices": "\u0645\u0643\u0627\u062a\u0628 \u0628\u0627\u0644\u0645\u062d\u0627\u0641\u0638",
    "up_bnd": "\u0627\u0631\u0641\u0639 \u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0628\u0627\u0642\u0627\u062a",
    "bnd": "\u0627\u0644\u0628\u0627\u0642\u0627\u062a",
    "bnd_info": "\u0627\u0644\u0628\u0627\u0642\u0627\u062a: \u0639\u062f\u062f \u0627\u0644\u0635\u0641\u0648\u0641",
    "bnd_unknown": "\u0628\u0627\u0642\u0627\u062a \u063a\u064a\u0631 \u0645\u0639\u0631\u0648\u0641\u0629",
    "bnd_nodist": "\u062e\u0637\u0648\u0637 \u063a\u064a\u0631 \u0645\u0648\u0632\u0639\u0629",
    "bnd_mm": "\u0628\u0627\u0642\u0627\u062a \u0645\u062e\u0627\u0644\u0641\u0629",
    "bnd_col": "\u0639\u0645\u0648\u062f \u0627\u0633\u0645 \u0627\u0644\u0628\u0627\u0642\u0629",
    "bnd_noname": "\u0645\u0627 \u0644\u0642\u064a\u062a \u0639\u0645\u0648\u062f \u0627\u0633\u0645 \u0627\u0644\u0628\u0627\u0642\u0629 \u0628\u062a\u0642\u0631\u064a\u0631 \u0627\u0644\u0628\u0627\u0642\u0627\u062a",
    "bnd_unames": "\u0623\u0633\u0645\u0627\u0621 \u0628\u0627\u0642\u0627\u062a \u063a\u064a\u0631 \u0645\u0639\u0631\u0648\u0641\u0629",
    "mm_ru": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0639\u0644\u0649 \u062e\u0637 \u0631\u064a\u062f",
    "mm_ur": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f \u0639\u0644\u0649 \u062e\u0637 \u064a\u0648\u0632",
    "mm_title": "\u0627\u0644\u0628\u0627\u0642\u0627\u062a \u0627\u0644\u0645\u062e\u0627\u0644\u0641\u0629 (\u0645\u0641\u0631\u0648\u0632\u0629 \u0644\u062d\u0627\u0644\u0647\u0627)",
    "line_lbl": "\u0631\u0642\u0645 \u0627\u0644\u062e\u0637",
    "dist_type": "\u0646\u0648\u0639 \u0627\u0644\u062a\u0648\u0632\u064a\u0639",
    "bundle_lbl": "\u0627\u0644\u0628\u0627\u0642\u0629",
    "mm_sheet": "\u0628\u0627\u0642\u0627\u062a \u0645\u062e\u0627\u0644\u0641\u0629",
    "bnd_empty": "\u0645\u0627 \u0644\u0642\u064a\u062a \u0628\u0627\u0642\u0627\u062a \u0645\u0637\u0627\u0628\u0642\u0629 \u0644\u062e\u0637\u0648\u0637 \u0645\u0648\u0632\u0639\u0629",
    "bnd_nomatch": "\u063a\u064a\u0631 \u0645\u0637\u0627\u0628\u0642\u0629 \u0644\u0644\u0648\u0643\u0644\u0627\u0621",
    "remain": "\u0628\u0627\u0642\u064a",
    "brand": "\u0628\u0631\u064a\u0642 \u0627\u0644\u0627\u0646\u0648\u0627\u0631",
    "rep": "\u0627\u0633\u0645 \u0627\u0644\u0645\u0646\u062f\u0648\u0628",
    "loading": "\u062c\u0627\u0631\u064a \u0642\u0631\u0627\u0621\u0629 \u0627\u0644\u0645\u0644\u0641 (\u0645\u0631\u0629 \u0648\u062d\u062f\u0629 \u0628\u0633)",
    "pts_sheet": "\u0627\u0644\u0646\u0642\u0627\u0637 \u0627\u0644\u0643\u0644\u064a\u0629",
    "pts_title": "\u062c\u062f\u0648\u0644 \u0645\u062a\u0627\u0628\u0639\u0629 \u0627\u0644\u0646\u0642\u0627\u0637 \u0648\u0627\u0644\u062e\u0637\u0648\u0637 - \u0634\u0647\u0631",
    "col_date": "\u0627\u0644\u062a\u0627\u0631\u064a\u062e",
    "h_net_pts": "\u062e\u0637 \u0646\u062a \u0646\u0642\u0627\u0637",
    "h_red_pts": "\u062e\u0637 \u0631\u064a\u062f \u0646\u0642\u0627\u0637",
    "h_use_pts": "\u062e\u0637 \u064a\u0648\u0632 \u0646\u0642\u0627\u0637",
    "h_net_cnt": "\u0639\u062f\u062f \u062e\u0637 \u0646\u062a",
    "h_red_cnt": "\u0639\u062f\u062f \u062e\u0637 \u0631\u064a\u062f",
    "h_use_cnt": "\u0639\u062f\u062f \u062e\u0637 \u064a\u0648\u0632",
    "f_daily": "\u0645\u062c\u0645\u0648\u0639 \u0627\u0644\u064a\u0648\u0645\u064a",
    "f_total": "\u0627\u0644\u0645\u062c\u0645\u0648\u0639 \u0627\u0644\u0643\u0644\u064a",
    "f_target": "\u0627\u0644\u062a\u0627\u0631\u0643\u062a",
    "f_ratio": "\u0627\u0644\u0646\u0633\u0628\u0629",
    "tgt_title": "\u0627\u0644\u062a\u0627\u0631\u0643\u062a \u0627\u0644\u0634\u0647\u0631\u064a (\u0646\u0642\u0627\u0637)",
    "tgt_net": "\u062a\u0627\u0631\u0643\u062a \u0646\u062a",
    "tgt_red": "\u062a\u0627\u0631\u0643\u062a \u0631\u064a\u062f",
    "tgt_use": "\u062a\u0627\u0631\u0643\u062a \u064a\u0648\u0632",
    "tgt_nosave": "\u0645\u0627 \u0643\u062f\u0631\u062a \u0623\u062d\u0641\u0638 \u0627\u0644\u062a\u0627\u0631\u0643\u062a \u0639\u0644\u0649 \u0627\u0644\u0633\u064a\u0631\u0641\u0631\u060c \u064a\u0628\u0642\u0649 \u0644\u0644\u062c\u0644\u0633\u0629 \u0627\u0644\u062d\u0627\u0644\u064a\u0629 \u0628\u0633",
    "key_detect": "\u062a\u0645 \u0627\u0643\u062a\u0634\u0627\u0641 \u0623\u0639\u0645\u062f\u0629 \u0631\u0642\u0645 \u0627\u0644\u0645\u0643\u062a\u0628 \u062a\u0644\u0642\u0627\u0626\u064a\u0627\u064b \u0644\u0623\u0646 B \u0648 F \u0645\u0627 \u062a\u0637\u0627\u0628\u0642\u0648\u0627",
    "key_none": "\u0645\u0627 \u0641\u064a \u0623\u064a \u0631\u0642\u0645 \u0645\u0643\u062a\u0628 \u0645\u0634\u062a\u0631\u0643 \u0628\u064a\u0646 \u0627\u0644\u0648\u0643\u0644\u0627\u0621 \u0648\u0627\u0644\u062a\u0648\u0632\u064a\u0639",
    "read_fail": "\u062a\u0639\u0630\u0631 \u0642\u0631\u0627\u0621\u0629 \u0627\u0644\u0645\u0644\u0641",
    "preview": "\u0645\u0639\u0627\u064a\u0646\u0629 \u0623\u0648\u0644 5 \u0635\u0641\u0648\u0641",
    "err": "\u062d\u062f\u062b \u062e\u0637\u0623",
    "monthly_sheet": "\u0634\u064a\u062a \u0627\u0644\u062a\u0648\u0632\u064a\u0639 \u0627\u0644\u0634\u0647\u0631\u064a",
}

RED = "\u0631\u064a\u062f"
USE = "\u064a\u0648\u0632"
MONTHLY = "\u0646\u062a \u0634\u0647\u0631\u064a"
DAILY = "\u0646\u062a \u064a\u0648\u0645\u064a"
BIG_ROUTER = "\u0631\u0627\u0648\u062a\u0631 \u0643\u0628\u064a\u0631"
SMALL_MIFI = "\u0631\u0627\u0648\u062a\u0631 \u0635\u063a\u064a\u0631"
USE_1M = "\u062e\u0637 \u064a\u0648\u0632 \u0645\u0644\u064a\u0648\u0646\u064a"
BLANK_LBL = "\u0628\u0644\u0627\u0646\u0643"

item_map = {
    "pre-paid sim card - type2": RED,
    "pre-paid sim card": RED,
    "youth line sim card": USE,
    "white list usim preloaded with iqd 40000": MONTHLY,
    "gprs sim card with package": DAILY,
    "gprs sim card with tozed cpe cat6": BIG_ROUTER,
    "gprs sim card with mifi huawei cat4": SMALL_MIFI,
    "nms-1 million coin silver": USE_1M,
}

# ===== monthly distribution sheet settings =====
M_ITEMS = [RED, USE, USE_1M, DAILY, MONTHLY, BIG_ROUTER, SMALL_MIFI]
M_STATUSES = ["distributed", "activated", "active"]

QUALITY_TYPE = {"2": "net", "93": "red", "4": "red", "32": "use"}

Q_TYPE_COL = 1
Q_OFFICE_COL = 2
DATE_COL = 3
Q_POINTS_COL = 17   # R = points
INC_OFFICE_COL = 0  # incentive: A = office number
INC_VALUE_COL = 7   # incentive: H = value (zeros removed)
WAL_OFFICE_COL = 1  # wallets: B = office number
WAL_FLAG_COL = 2    # wallets: C = flag (keep rows with Y)

# ===== Bundles report (A = office number, H = sub-type) =====
BND_OFFICE_COL = 0   # bundles: A = office number (matched with agents B)
BND_SUB_COL = 7

BUNDLE_DIRECT = {
    "red 15,000 monthly package": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 15",
    "red 15 4 weeks": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 3 \u0623\u0634\u0647\u0631",
    "red 30 4 weeks": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f \u0634\u0647\u0631\u064a 30",
    "monthly10gb": "\u0628\u0627\u0642\u0629 \u0625\u0646\u062a\u0631\u0646\u062a \u0634\u0647\u0631\u064a\u0629 \u0628\u0644\u0627 \u062d\u062f\u0648\u062f",
    "yooz 15-extra voice for 12 weeks": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "yooz 15 extra voice": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "yooz 15 mix": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "yooz 15-extra data for 12 weeks": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "yooz 15 extra data": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "yooz 30 extra data": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 30",
    "yooz 25 extra data": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 25",
    "yooz 25 mix": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 25",
    "yooz 5 extra voice": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0627\u0645 5",
    "yooz 5 extra data": "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0627\u0645 5",
    "red 5,000 weekly package": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 5",
    "red 10,000 for 20 days": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 10",
    "red 25 4 weeks": "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 25",
}

BUNDLE_SPLIT = {
    "unlimited weekly 4g": {
        "base": "\u0628\u0627\u0642\u0629 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000 \u0628\u0644\u0627 \u062d\u062f\u0648\u062f",
        "sub": [
            ("usim-gprs-pkg", "\u0628\u0627\u0642\u0629 \u0646\u062a \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000"),
            ("usim-prp2", "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 12000 \u0623\u0633\u0628\u0648\u0639\u064a\u0629"),
            ("usim-yout", "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000"),
        ],
    },
    "unlimited 4g bundle": {
        "base": "\u0628\u0627\u0642\u0629 \u0623\u0631\u0628\u0639\u064a\u0646",
        "sub": [
            ("usim-gprs-pkg", "\u0628\u0627\u0642\u0629 \u0646\u062a \u0623\u0633\u0628\u0648\u0639\u064a\u0629 40"),
            ("usim-prp2", "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 40 \u0623\u0633\u0628\u0648\u0639\u064a\u0629"),
            ("usim-yout", "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 40"),
            ("usim-40kv", "\u0628\u0627\u0642\u0629 \u0646\u062a \u0634\u0647\u0631\u064a 40"),
        ],
    },
    "unlimited premium package fup": {
        "base": "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645",
        "sub": [
            ("usim-gprs-rtr2", "\u0628\u0627\u0642\u0629 \u0631\u0627\u0648\u062a\u0631 110"),
            ("usim-40kv", "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0646\u062a \u0623\u0628\u0648 40 \u0628\u062a\u0639\u0628\u0626\u0629 50"),
            ("usim-prp2", "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645 \u0631\u064a\u062f"),
            ("usim-prp", "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645 \u0631\u064a\u062f"),
        ],
    },
}

BUNDLE_ORDER = [
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 5",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 10",
    "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 12000 \u0623\u0633\u0628\u0648\u0639\u064a\u0629",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 15",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 3 \u0623\u0634\u0647\u0631",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 25",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f \u0634\u0647\u0631\u064a 30",
    "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 40 \u0623\u0633\u0628\u0648\u0639\u064a\u0629",
    "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645 \u0631\u064a\u062f",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0627\u0645 5",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 25",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 30",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 40",
    "\u0628\u0627\u0642\u0629 \u0646\u062a \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000",
    "\u0628\u0627\u0642\u0629 \u0646\u062a \u0623\u0633\u0628\u0648\u0639\u064a\u0629 40",
    "\u0628\u0627\u0642\u0629 \u0646\u062a \u0634\u0647\u0631\u064a 40",
    "\u0628\u0627\u0642\u0629 \u0625\u0646\u062a\u0631\u0646\u062a \u0634\u0647\u0631\u064a\u0629 \u0628\u0644\u0627 \u062d\u062f\u0648\u062f",
    "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0646\u062a \u0623\u0628\u0648 40 \u0628\u062a\u0639\u0628\u0626\u0629 50",
    "\u0628\u0627\u0642\u0629 \u0631\u0627\u0648\u062a\u0631 110",
    "\u0628\u0627\u0642\u0629 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000 \u0628\u0644\u0627 \u062d\u062f\u0648\u062f",
    "\u0628\u0627\u0642\u0629 \u0623\u0631\u0628\u0639\u064a\u0646",
    "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645",
]

BUNDLE_FAMILY = {}
for _lbl in [
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 3 \u0623\u0634\u0647\u0631",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f \u0634\u0647\u0631\u064a 30",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 5",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 10",
    "\u0628\u0627\u0642\u0629 \u0631\u064a\u062f 25",
    "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 12000 \u0623\u0633\u0628\u0648\u0639\u064a\u0629",
    "\u0628\u0627\u0642\u0629 \u062e\u0637 \u0631\u064a\u062f 40 \u0623\u0633\u0628\u0648\u0639\u064a\u0629",
    "\u0628\u0627\u0642\u0629 \u0628\u0631\u064a\u0645\u064a\u0648\u0645 \u0631\u064a\u062f",
]:
    BUNDLE_FAMILY[_lbl] = "red"
for _lbl in [
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0627\u0645 5",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 15000",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 25",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 30",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 12000",
    "\u0628\u0627\u0642\u0629 \u064a\u0648\u0632 \u0623\u0633\u0628\u0648\u0639\u064a\u0629 40",
]:
    BUNDLE_FAMILY[_lbl] = "use"

LINE_FAMILY = {RED: "red", USE: "use", USE_1M: "use"}

MONTHS = [
    "\u0643\u0627\u0646\u0648\u0646 \u0627\u0644\u062b\u0627\u0646\u064a",
    "\u0634\u0628\u0627\u0637",
    "\u0622\u0630\u0627\u0631",
    "\u0646\u064a\u0633\u0627\u0646",
    "\u0623\u064a\u0627\u0631",
    "\u062d\u0632\u064a\u0631\u0627\u0646",
    "\u062a\u0645\u0648\u0632",
    "\u0622\u0628",
    "\u0623\u064a\u0644\u0648\u0644",
    "\u062a\u0634\u0631\u064a\u0646 \u0627\u0644\u0623\u0648\u0644",
    "\u062a\u0634\u0631\u064a\u0646 \u0627\u0644\u062b\u0627\u0646\u064a",
    "\u0643\u0627\u0646\u0648\u0646 \u0627\u0644\u0623\u0648\u0644",
]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def last10(num):
    if pd.isna(num):
        return None
    s = str(num).strip()
    if re.fullmatch(r"\d+\.0+", s):
        s = s.split(".")[0]
    elif re.fullmatch(r"\d+(\.\d+)?[eE]\+?\d+", s):
        try:
            s = str(int(float(s)))
        except Exception:
            pass
    s = re.sub(r"\D", "", s)
    if s:
        return s[-10:]
    return None


def letter_of(i):
    if i < 26:
        return chr(65 + i)
    return str(i)


def text_columns(df, key_idx):
    # columns (among the first 8) that mostly hold names (contain letters), without the key column
    found = []
    for i in range(min(df.shape[1], 8)):
        if i == key_idx:
            continue
        vals = df.iloc[:, i].dropna().astype(str).head(50)
        if len(vals) == 0:
            continue
        with_letters = vals.str.contains(r"[A-Za-z\u0600-\u06FF]", regex=True).sum()
        if with_letters >= 0.6 * len(vals):
            found.append(i)
    return found


def find_key_columns(agents, dist):
    # fallback when B (agents) and F (distribution) share no numbers:
    # find the two columns that share the most office numbers
    a_sets = {}
    for i in range(min(agents.shape[1], 10)):
        a_sets[i] = set(agents.iloc[:, i].apply(last10).dropna())
    d_sets = {}
    for j in range(min(dist.shape[1], 12)):
        d_sets[j] = set(dist.iloc[:, j].apply(last10).dropna())
    best = (0, 1, 5)
    for i, a_set in a_sets.items():
        if len(a_set) == 0:
            continue
        for j, d_set in d_sets.items():
            n = len(a_set & d_set)
            if n > best[0]:
                best = (n, i, j)
    return best


def clean_code(val):
    if pd.isna(val):
        return None
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def detect_kind(raw):
    head = raw[:2048]
    if head[:4] == b"PK\x03\x04":
        return "xlsx"
    if head[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        return "xls"
    low = head.lower()
    if b"<html" in low or b"<table" in low or b"<!doctype" in low:
        return "html"
    return "csv"


def decode_text(raw):
    if raw[:3] == b"\xef\xbb\xbf":
        return raw.decode("utf-8-sig", errors="ignore")
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16", errors="ignore")
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError:
        return raw.decode("cp1256", errors="ignore")


def read_any(raw, kind):
    if kind == "xlsx":
        return pd.read_excel(BytesIO(raw), sheet_name=None, dtype=str, engine="openpyxl")
    if kind == "xls":
        return pd.read_excel(BytesIO(raw), sheet_name=None, dtype=str, engine="xlrd")
    text = decode_text(raw)
    if kind == "html":
        tables = pd.read_html(StringIO(text))
        return {"sheet1": tables[0].astype(str)}
    head = text[:4096]
    sep = ","
    if head.count("\t") > head.count(","):
        sep = "\t"
    return {"sheet1": pd.read_csv(StringIO(text), sep=sep, dtype=str)}


@st.cache_data(show_spinner=T["loading"], max_entries=16)
def read_cached(raw, kind):
    # each file is parsed once; reruns of the page reuse the saved result
    return read_any(raw, kind)


def load_report(file, label):
    raw = file.getvalue()
    kind = detect_kind(raw)
    try:
        return read_cached(raw, kind)
    except Exception as e:
        st.error(T["read_fail"] + " (" + label + ", " + kind + "): " + str(e))
        return None


def norm_name(x):
    if pd.isna(x):
        return ""
    return re.sub(r"\s+", " ", str(x)).strip().lower()


def classify_bundle(name_n, sub_n):
    if name_n in BUNDLE_SPLIT:
        spec = BUNDLE_SPLIT[name_n]
        for key, label in spec["sub"]:
            if key in sub_n:
                return label
        return spec["base"]
    return BUNDLE_DIRECT.get(name_n)


def detect_bundle_col(b):
    known = set(BUNDLE_DIRECT) | set(BUNDLE_SPLIT)
    best = None
    best_score = 0
    for i in range(b.shape[1]):
        score = int(b.iloc[:, i].map(norm_name).isin(known).sum())
        if score > best_score:
            best = i
            best_score = score
    return best, best_score


def pick_sheet(sheets, label):
    names = list(sheets.keys())
    if len(names) == 1:
        return sheets[names[0]]
    chosen = st.selectbox(T["sheet"] + " " + label, names, key="sheet_" + label)
    return sheets[chosen]


def compute(agents, dist, q, inc=None, wal=None, bnd=None):
    a_idx = 1   # agents: B = office number
    d_idx = 5   # distribution: F = office number
    a_keys = agents.iloc[:, a_idx].apply(last10)
    d_keys = dist.iloc[:, d_idx].apply(last10)
    if len(set(a_keys.dropna()) & set(d_keys.dropna())) == 0:
        n_best, bi, bj = find_key_columns(agents, dist)
        if n_best > 0:
            a_idx = bi
            d_idx = bj
            a_keys = agents.iloc[:, a_idx].apply(last10)
            st.warning(
                T["key_detect"] + ": " + T["agents"] + " = " + letter_of(a_idx)
                + " | " + T["dist"] + " = " + letter_of(d_idx)
            )
        else:
            st.warning(
                T["key_none"]
                + " | " + T["agents"] + " (B): " + str(agents.iloc[:3, 1].tolist()) + " -> " + str(a_keys.head(3).tolist())
                + " | " + T["dist"] + " (F): " + str(dist.iloc[:3, 5].tolist()) + " -> " + str(d_keys.head(3).tolist())
            )

    if a_idx == 1:
        name_idx = 0
        rep_idx = 2
    else:
        txt = text_columns(agents, a_idx)
        if len(txt) > 0:
            name_idx = txt[0]
        else:
            name_idx = 0 if a_idx != 0 else 1
        if len(txt) > 1:
            rep_idx = txt[1]
        else:
            rep_idx = None
    if rep_idx is not None and agents.shape[1] > rep_idx:
        rep_vals = agents.iloc[:, rep_idx]
    else:
        rep_vals = ""
    agents_df = pd.DataFrame({
        "key": a_keys,
        T["office_name"]: agents.iloc[:, name_idx],
        T["rep"]: rep_vals,
    })
    agents_df = agents_df.dropna(subset=["key"]).drop_duplicates("key")

    d = pd.DataFrame({
        "line": dist.iloc[:, 0].astype(str).str.strip(),
        "item": dist.iloc[:, 2].astype(str).str.strip().str.lower(),
        "status": dist.iloc[:, 4].astype(str).str.strip().str.lower(),
        "key": dist.iloc[:, d_idx].apply(last10),
        "date": pd.to_datetime(dist.iloc[:, 8], errors="coerce"),
    })

    d_raw = d.copy()  # raw rows, used by the monthly sheet (its own de-duplication)
    total_rows = len(d)
    bad_dates = int(d["date"].isna().sum())

    d = d.sort_values("date", kind="stable")
    d = d.drop_duplicates(subset="line", keep="last")
    after_dedup = len(d)

    d = d[d["status"] == "distributed"].copy()
    d[T["item"]] = d["item"].map(item_map)
    d = d.dropna(subset=[T["item"]])

    if d.empty:
        st.error(T["empty"])
        return None

    merged = d.merge(agents_df, on="key", how="left")
    unmatched = int(merged[T["office_name"]].isna().sum())

    result = pd.crosstab(
        [merged["key"], merged[T["office_name"]].fillna(T["unmatched"])],
        merged[T["item"]],
    )
    for col in [RED, USE, DAILY, MONTHLY, BIG_ROUTER, SMALL_MIFI]:
        if col not in result.columns:
            result[col] = 0
    result[T["total"]] = result[RED] + result[USE]
    result = result.sort_values(T["total"], ascending=False).reset_index()
    result = result.rename(columns={"key": T["office_no"]})

    q_msg = None
    points_src = None
    if q is not None:
        qd = pd.DataFrame({
            "group": q.iloc[:, Q_TYPE_COL].apply(clean_code).map(QUALITY_TYPE),
            "key": q.iloc[:, Q_OFFICE_COL].apply(last10),
            "date": pd.to_datetime(q.iloc[:, DATE_COL], errors="coerce"),
        })

        if q.shape[1] > Q_POINTS_COL:
            pts_raw = q.iloc[:, Q_POINTS_COL].astype(str).str.replace(",", "", regex=False)
            qd["pts"] = pd.to_numeric(pts_raw, errors="coerce")
        else:
            qd["pts"] = 0.0
        q_total = len(qd)
        q_ignored = int(qd["group"].isna().sum())
        q_baddate = int(qd["date"].isna().sum())

        points_src = qd.dropna(subset=["group"])[["group", "date", "pts"]].copy()
        qd = qd.dropna(subset=["group", "key"]).copy()

        tot = pd.crosstab(qd["key"], qd["group"])
        for g in ["red", "use", "net"]:
            if g not in tot.columns:
                tot[g] = 0
        tot = tot[["red", "use", "net"]]
        tot.columns = [T["act_red"], T["act_use"], T["act_net"]]

        ru_mask = qd["group"].isin(["red", "use"])
        pts_ru = qd[ru_mask].groupby("key")["pts"].sum()
        pts_net = qd[qd["group"] == "net"].groupby("key")["pts"].sum()
        pts_tbl = pd.DataFrame({T["pts_ru"]: pts_ru, T["pts_net"]: pts_net})
        tot = tot.join(pts_tbl, how="outer")

        qdd = qd.dropna(subset=["date"]).copy()
        if qdd.empty:
            last_day = None
            last_tbl = pd.DataFrame(columns=[T["last_ru"], T["last_net"], T["pts_last_ru"], T["pts_last_net"]])
        else:
            qdd["day"] = qdd["date"].dt.normalize()
            last_day = qdd["day"].max()
            ld = qdd[qdd["day"] == last_day]
            ru = ld[ld["group"].isin(["red", "use"])].groupby("key").size()
            nt = ld[ld["group"] == "net"].groupby("key").size()
            pr = ld[ld["group"].isin(["red", "use"])].groupby("key")["pts"].sum()
            pn = ld[ld["group"] == "net"].groupby("key")["pts"].sum()
            last_tbl = pd.DataFrame({
                T["last_ru"]: ru,
                T["last_net"]: nt,
                T["pts_last_ru"]: pr,
                T["pts_last_net"]: pn,
            })

        q_all = tot.join(last_tbl, how="outer").fillna(0)
        q_all.index.name = T["office_no"]
        q_all = q_all.reset_index()

        if last_day is not None:
            day_txt = str(last_day.date())
            q_all = q_all.rename(columns={
                T["last_ru"]: T["last_ru"] + " " + day_txt,
                T["last_net"]: T["last_net"] + " " + day_txt,
                T["pts_last_ru"]: T["pts_last_ru"] + " " + day_txt,
                T["pts_last_net"]: T["pts_last_net"] + " " + day_txt,
            })

        result = result.merge(q_all, on=T["office_no"], how="outer")

        name_map = dict(zip(agents_df["key"], agents_df[T["office_name"]]))
        result[T["office_name"]] = result[T["office_name"]].fillna(
            result[T["office_no"]].map(name_map)
        ).fillna(T["unmatched"])

        num_cols = [c for c in result.columns if c not in (T["office_no"], T["office_name"])]
        for c in num_cols:
            result[c] = result[c].fillna(0)
            if c.startswith(T["pts_word"]):
                result[c] = result[c].round(2)
            else:
                result[c] = result[c].astype(int)

        last_txt = "-"
        if last_day is not None:
            last_txt = str(last_day.date())
        q_msg = (
            T["q_info"] + ": " + str(q_total)
            + " | " + T["q_ignored"] + ": " + str(q_ignored)
            + " | " + T["q_baddate"] + ": " + str(q_baddate)
            + " | " + T["q_last"] + ": " + last_txt
        )

    inc_msg = None
    if inc is not None:
        inc_df = pd.DataFrame({
            "key": inc.iloc[:, INC_OFFICE_COL].apply(last10),
            "val": pd.to_numeric(
                inc.iloc[:, INC_VALUE_COL].astype(str).str.replace(",", "", regex=False),
                errors="coerce",
            ),
        })
        inc_total = len(inc_df)
        inc_df = inc_df.dropna(subset=["key", "val"])
        inc_df = inc_df[inc_df["val"] != 0]
        inc_after = len(inc_df)
        inc_keys = inc_df["key"].drop_duplicates()
        inc_tbl = pd.DataFrame({T["office_no"]: inc_keys, T["inc"]: "done"})

        result = result.merge(inc_tbl, on=T["office_no"], how="outer")

        name_map2 = dict(zip(agents_df["key"], agents_df[T["office_name"]]))
        result[T["office_name"]] = result[T["office_name"]].fillna(
            result[T["office_no"]].map(name_map2)
        ).fillna(T["unmatched"])
        result[T["inc"]] = result[T["inc"]].fillna("")

        skip_cols = (T["office_no"], T["office_name"], T["inc"])
        num_cols2 = [c for c in result.columns if c not in skip_cols]
        for c in num_cols2:
            result[c] = result[c].fillna(0)
            if c.startswith(T["pts_word"]):
                result[c] = result[c].round(2)
            else:
                result[c] = result[c].astype(int)

        inc_msg = (
            T["inc_info"] + ": " + str(inc_total)
            + " | " + T["inc_after"] + ": " + str(inc_after)
            + " | " + T["inc_offices"] + ": " + str(len(inc_keys))
        )

    wal_msg = None
    if wal is not None:
        wal_df = pd.DataFrame({
            "key": wal.iloc[:, WAL_OFFICE_COL].apply(last10),
            "flag": wal.iloc[:, WAL_FLAG_COL].astype(str).str.strip().str.lower(),
        })
        wal_total = len(wal_df)
        wal_df = wal_df.dropna(subset=["key"])
        wal_df = wal_df[wal_df["flag"] == "y"]
        wal_after = len(wal_df)
        wal_keys = wal_df["key"].drop_duplicates()
        wal_tbl = pd.DataFrame({T["office_no"]: wal_keys, T["wal"]: "done"})

        result = result.merge(wal_tbl, on=T["office_no"], how="outer")

        name_map3 = dict(zip(agents_df["key"], agents_df[T["office_name"]]))
        result[T["office_name"]] = result[T["office_name"]].fillna(
            result[T["office_no"]].map(name_map3)
        ).fillna(T["unmatched"])

        text_cols = [T["inc"], T["wal"]]
        for c in text_cols:
            if c in result.columns:
                result[c] = result[c].fillna("")

        skip_cols3 = [T["office_no"], T["office_name"]] + text_cols
        num_cols3 = [c for c in result.columns if c not in skip_cols3]
        for c in num_cols3:
            result[c] = result[c].fillna(0)
            if c.startswith(T["pts_word"]):
                result[c] = result[c].round(2)
            else:
                result[c] = result[c].astype(int)

        wal_msg = (
            T["wal_info"] + ": " + str(wal_total)
            + " | " + T["wal_after"] + ": " + str(wal_after)
            + " | " + T["wal_offices"] + ": " + str(len(wal_keys))
        )

    bnd_msg = None
    bnd_warn = None
    bundle_df = None
    mismatch_df = None
    if bnd is not None:
        name_idx, hits = detect_bundle_col(bnd)
        if name_idx is None:
            st.error(T["bnd_noname"])
        else:
            if bnd.shape[1] > BND_SUB_COL:
                sub_vals = bnd.iloc[:, BND_SUB_COL].map(norm_name)
            else:
                sub_vals = ""
            b = pd.DataFrame({
                "key": bnd.iloc[:, BND_OFFICE_COL].apply(last10),
                "name": bnd.iloc[:, name_idx].map(norm_name),
                "sub": sub_vals,
            })
            b_total = len(b)
            b["label"] = [classify_bundle(n, s2) for n, s2 in zip(b["name"], b["sub"])]
            unknown = b.loc[b["label"].isna() & (b["name"] != ""), "name"].value_counts().head(8)
            b_unknown = int(b["label"].isna().sum())
            b = b.dropna(subset=["label", "key"])

            agent_keys = set(agents_df["key"])
            b_nomatch = int((~b["key"].isin(agent_keys)).sum())

            if len(b) > 0:
                tbl = b.groupby(["key", "label"]).size().unstack(fill_value=0)
            else:
                tbl = pd.DataFrame()
            tbl.index.name = T["office_no"]
            tbl.columns.name = None
            tbl = tbl.reset_index()
            name_map5 = dict(zip(agents_df["key"], agents_df[T["office_name"]]))
            tbl.insert(0, T["office_name"], tbl[T["office_no"]].map(name_map5).fillna(T["unmatched"]))
            rep_map5 = dict(zip(agents_df["key"], agents_df[T["rep"]]))
            tbl.insert(2, T["rep"], tbl[T["office_no"]].map(rep_map5).fillna(""))
            ordered_b = [c for c in BUNDLE_ORDER if c in tbl.columns]
            tbl = tbl[[T["office_name"], T["office_no"], T["rep"]] + ordered_b]
            if len(ordered_b) > 0:
                tot_b = tbl[ordered_b].sum(axis=1)
                tbl = tbl.assign(_t=tot_b).sort_values("_t", ascending=False).drop(columns="_t")
            bundle_df = tbl.reset_index(drop=True)

            if hits == 0:
                bnd_warn = T["bnd_noname"]
            if len(unknown) > 0:
                names_txt = " ; ".join(str(k) + " (" + str(v) + ")" for k, v in unknown.items())
                bnd_warn = T["bnd_unames"] + ": " + names_txt
            col_letter = chr(65 + name_idx) if name_idx < 26 else str(name_idx)
            bnd_msg = (
                T["bnd_info"] + ": " + str(b_total)
                + " | " + T["bnd_unknown"] + ": " + str(b_unknown)
                + " | " + T["bnd_nomatch"] + ": " + str(b_nomatch)
                + " | " + T["bnd_col"] + ": " + col_letter
            )

    # ===== monthly distribution sheet =====
    # month = month of the latest date in the distribution file.
    # inside that month: same line (column A) repeated -> keep only the latest date,
    # then keep statuses distributed + active, ALL item types (unknown ones keep their own name)
    monthly_df = None
    monthly_month = None
    monthly_msg = None
    dd = d_raw.dropna(subset=["date"])
    if not dd.empty:
        last_d = dd["date"].max()
        monthly_month = int(last_d.month)
        in_m = dd[(dd["date"].dt.year == last_d.year) & (dd["date"].dt.month == last_d.month)].copy()
        in_m["lkey"] = in_m["line"].map(lambda s_: last10(s_) or s_)
        in_m = in_m.sort_values("date", kind="stable")
        in_m = in_m.drop_duplicates(subset="lkey", keep="last")
        m_rows_month = len(in_m)
        in_m = in_m[in_m["status"].isin(M_STATUSES)].copy()
        m_rows_status = len(in_m)
        m_no_key = int(in_m["key"].isna().sum())
        monthly_msg = (
            "monthly: rows(no dup)=" + str(m_rows_month)
            + " | after status=" + str(m_rows_status)
            + " | no office number=" + str(m_no_key)
            + " | bad dates=" + str(bad_dates)
        )
        raw_item = in_m["item"].replace({"": BLANK_LBL, "nan": BLANK_LBL, "none": BLANK_LBL})
        in_m[T["item"]] = in_m["item"].map(item_map).fillna(raw_item)
        in_m = in_m.dropna(subset=["key"])
        if not in_m.empty:
            mt = pd.crosstab(in_m["key"], in_m[T["item"]])
            for col in M_ITEMS:
                if col not in mt.columns:
                    mt[col] = 0
            extra_items = [c for c in mt.columns if c not in M_ITEMS]
            mt = mt[M_ITEMS + extra_items]
            mt.columns.name = None
            mt[T["total"]] = mt.sum(axis=1)
            mt = mt.sort_values(T["total"], ascending=False)
            mt.index.name = T["office_no"]
            mt = mt.reset_index()
            m_names = dict(zip(agents_df["key"], agents_df[T["office_name"]]))
            m_reps = dict(zip(agents_df["key"], agents_df[T["rep"]]))
            mt.insert(0, T["office_name"], mt[T["office_no"]].map(m_names).fillna(T["unmatched"]))
            mt.insert(2, T["rep"], mt[T["office_no"]].map(m_reps).fillna("").astype(str))
            monthly_df = mt

    rep_map = dict(zip(agents_df["key"], agents_df[T["rep"]]))
    result[T["rep"]] = result[T["office_no"]].map(rep_map).fillna("")
    order_base = [
        T["office_name"], T["office_no"], T["rep"],
        RED, USE, DAILY, MONTHLY, BIG_ROUTER, SMALL_MIFI,
        T["total"],
        T["act_red"], T["act_use"], T["act_net"],
        T["last_ru"], T["last_net"],
        T["pts_ru"], T["pts_net"], T["pts_last_ru"], T["pts_last_net"],
        T["inc"], T["wal"],
    ]
    ordered = []
    for base in order_base:
        for c in result.columns:
            if c == base or str(c).startswith(base + " "):
                if c not in ordered:
                    ordered.append(c)
    rest = [c for c in result.columns if c not in ordered]
    result = result[ordered + rest]

    ren = {}
    for c in [RED, USE, DAILY, MONTHLY, BIG_ROUTER, SMALL_MIFI, USE_1M]:
        if c in result.columns:
            ren[c] = T["remain"] + " " + c
    result = result.rename(columns=ren)

    st.success(T["done"] + ": " + str(len(result)) + " | " + T["unm_rows"] + ": " + str(unmatched))
    st.info(
        T["orig"] + ": " + str(total_rows)
        + " | " + T["dedup"] + ": " + str(after_dedup)
        + " | " + T["real"] + ": " + str(len(d))
        + " | " + T["bad_dates"] + ": " + str(bad_dates)
    )
    if q_msg is not None:
        st.info(q_msg)
    if inc_msg is not None:
        st.info(inc_msg)
    if wal_msg is not None:
        st.info(wal_msg)
    if bnd_msg is not None:
        st.info(bnd_msg)
    if bnd_warn is not None:
        st.warning(bnd_warn)
    result.attrs["mismatch"] = mismatch_df
    result.attrs["bundles"] = bundle_df
    result.attrs["points_src"] = points_src
    result.attrs["monthly"] = monthly_df
    result.attrs["monthly_month"] = monthly_month
    result.attrs["monthly_msg"] = monthly_msg
    return result


# ===== formatting settings (change here) =====
FONT_SIZE = 14
POINTS_SHEET_RTL = False   # False = points sheet starts from the left side (True = from the right)
HEADER_LINE_CHARS = 19   # approx. characters per line in a header cell
HEADER_LINE_PT = 19      # height of one text line (points)
COL_WIDTH = 24
GRAY_BODY = "D9D9D9"
GRAY_HEADER = "A6A6A6"
YELLOW = "FFFF00"


def fmt_num(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        v = int(v)
    if isinstance(v, int):
        return "{:,}".format(v)
    return "{:,.2f}".format(v)


def build_points_table(src, tgt_net, tgt_red, tgt_use):
    # daily table of points and counts per type (net / red / use) for the month
    s2 = src.copy()
    s2["pts"] = s2["pts"].fillna(0)
    dated = s2.dropna(subset=["date"]).copy()
    if dated.empty:
        return None
    dated["day"] = dated["date"].dt.normalize()
    last_day = dated["day"].max()
    year = last_day.year
    month = last_day.month
    ndays = calendar.monthrange(year, month)[1]
    month_name = MONTHS[month - 1]
    in_month = dated[(dated["day"].dt.year == year) & (dated["day"].dt.month == month)]
    groups = ["net", "red", "use"]

    rows = []
    for day in range(1, ndays + 1):
        label = month_name + " " + str(day).zfill(2)
        if day > last_day.day:
            rows.append([label, None, None, None, None, None, None])
            continue
        sub = in_month[in_month["day"].dt.day == day]
        vals = []
        for g in groups:
            vals.append(float(sub.loc[sub["group"] == g, "pts"].sum()))
        for g in groups:
            vals.append(int((sub["group"] == g).sum()))
        rows.append([label] + vals)

    daily_sum = []
    for i in range(1, 7):
        daily_sum.append(sum(r[i] for r in rows if r[i] is not None))

    grand = []
    for g in groups:
        grand.append(float(s2.loc[s2["group"] == g, "pts"].sum()))
    for g in groups:
        grand.append(int((s2["group"] == g).sum()))

    target = [tgt_net, tgt_red, tgt_use]
    ratio = []
    for i in range(3):
        if target[i]:
            ratio.append(grand[i] / target[i])
        else:
            ratio.append(None)

    heads = [T["col_date"], T["h_net_pts"], T["h_red_pts"], T["h_use_pts"],
             T["h_net_cnt"], T["h_red_cnt"], T["h_use_cnt"]]
    disp = []
    for r in rows:
        disp.append([r[0]] + [fmt_num(v) for v in r[1:]])
    disp.append([T["f_daily"]] + [fmt_num(v) for v in daily_sum])
    disp.append([T["f_total"]] + [fmt_num(v) for v in grand])
    disp.append([T["f_target"]] + [fmt_num(v) for v in target] + ["", "", ""])
    pct = []
    for v in ratio:
        if v is None:
            pct.append("")
        else:
            pct.append("{:.2f}%".format(v * 100))
    disp.append([T["f_ratio"]] + pct + ["", "", ""])
    display_df = pd.DataFrame(disp, columns=heads)

    return {
        "title": T["pts_title"] + " " + month_name,
        "heads": heads,
        "rows": rows,
        "daily_sum": daily_sum,
        "grand": grand,
        "target": target,
        "ratio": ratio,
        "display": display_df,
    }


def write_points_sheet(ws, info):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    blue = "1F4E79"
    ws.sheet_view.rightToLeft = POINTS_SHEET_RTL
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    light = Side(style="thin", color="BFBFBF")
    box = Border(left=light, right=light, top=light, bottom=light)
    head_fill = PatternFill("solid", start_color=blue, end_color=blue)
    band_fill = PatternFill("solid", start_color="EEF3FA", end_color="EEF3FA")
    white_fill = PatternFill("solid", start_color="FFFFFF", end_color="FFFFFF")
    foot_fill = PatternFill("solid", start_color="DCE6F1", end_color="DCE6F1")
    ncols = 7

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    t = ws.cell(row=1, column=1, value=info["title"])
    t.font = Font(size=FONT_SIZE + 4, bold=True, color=blue)
    t.alignment = center
    ws.row_dimensions[1].height = 34

    hr = 3
    for j, h in enumerate(info["heads"], start=1):
        c = ws.cell(row=hr, column=j, value=h)
        c.font = Font(size=FONT_SIZE, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = center
        c.border = box
    ws.row_dimensions[hr].height = 30

    r = hr + 1
    first_data = r
    for k, row in enumerate(info["rows"]):
        fill = band_fill if k % 2 == 0 else white_fill
        for j, v in enumerate(row, start=1):
            if isinstance(v, float) and v == int(v):
                v = int(v)
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(size=FONT_SIZE)
            c.fill = fill
            c.alignment = center
            c.border = box
            if j > 1:
                c.number_format = "#,##0"
        r += 1
    last_data = r - 1

    def footer(label, values, number_format, last=False):
        nonlocal r
        cells = [label] + list(values)
        while len(cells) < ncols:
            cells.append(None)
        for j, v in enumerate(cells, start=1):
            if isinstance(v, float) and number_format == "#,##0" and v == int(v):
                v = int(v)
            c = ws.cell(row=r, column=j, value=v)
            c.font = Font(size=FONT_SIZE, bold=True)
            c.fill = foot_fill
            c.alignment = center
            if last:
                dbl = Side(style="double", color=blue)
                c.border = Border(left=light, right=light, top=light, bottom=dbl)
            else:
                c.border = box
            if j > 1:
                c.number_format = number_format
        r += 1

    footer(T["f_daily"], info["daily_sum"], "#,##0")
    footer(T["f_total"], info["grand"], "#,##0")
    footer(T["f_target"], info["target"], "#,##0")
    footer(T["f_ratio"], info["ratio"], "0.00%", last=True)

    for i in range(1, ncols + 1):
        ws.column_dimensions[get_column_letter(i)].width = 20
    ws.auto_filter.ref = "A" + str(hr) + ":" + get_column_letter(ncols) + str(last_data)
    ws.freeze_panes = "A" + str(hr + 1)


def format_excel(df, extras=None, points=None, after_points=None):
    # writes df (and optional extra sheet) to xlsx bytes: gray cells, yellow first column,
    # borders, centered text, equal column widths, font size 14
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    gray_body = PatternFill("solid", start_color=GRAY_BODY, end_color=GRAY_BODY)
    gray_head = PatternFill("solid", start_color=GRAY_HEADER, end_color=GRAY_HEADER)
    yellow = PatternFill("solid", start_color=YELLOW, end_color=YELLOW)

    def style_sheet(ws):
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                is_header = cell.row == 1
                cell.font = Font(size=FONT_SIZE, bold=is_header)
                cell.alignment = center
                cell.border = border
                if cell.column == 1:
                    cell.fill = yellow
                elif is_header:
                    cell.fill = gray_head
                else:
                    cell.fill = gray_body
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = COL_WIDTH
        max_lines = 1
        for cell in ws[1]:
            if cell.value is not None:
                n = len(textwrap.wrap(str(cell.value), HEADER_LINE_CHARS)) or 1
                max_lines = max(max_lines, n)
        ws.row_dimensions[1].height = max_lines * HEADER_LINE_PT + 6
        ws.freeze_panes = "B2"

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sheet1")
        style_sheet(writer.sheets["Sheet1"])
        if extras is not None:
            for sheet_name, sheet_df in extras:
                sheet_df.to_excel(writer, index=False, sheet_name=sheet_name)
                style_sheet(writer.sheets[sheet_name])
        if points is not None:
            ws_points = writer.book.create_sheet(T["pts_sheet"])
            write_points_sheet(ws_points, points)
        if after_points is not None:
            for sheet_name, sheet_df in after_points:
                sheet_df.to_excel(writer, index=False, sheet_name=sheet_name)
                style_sheet(writer.sheets[sheet_name])
    return buf.getvalue()


def style_screen(df):
    # on-screen version of the same look (colors + centered text)
    first = df.columns[0]
    sty = df.style.set_properties(**{
        "background-color": "#D9D9D9",
        "color": "black",
        "text-align": "center",
        "font-size": "14px",
        "border": "1px solid black",
    })
    sty = sty.set_properties(subset=[first], **{"background-color": "#FFFF00"})
    return sty


try:
    TARGETS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "targets.json")
except NameError:
    TARGETS_FILE = "targets.json"


def target_month_key(now=None):
    # month of the report date (= yesterday, Iraq time).
    # the target starts again from 0 when this month changes
    if now is None:
        now = datetime.now(timezone(timedelta(hours=3)))
    day = now - timedelta(days=1)
    return str(day.year) + "-" + str(day.month).zfill(2)


def load_targets():
    try:
        with open(TARGETS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_targets(month_key, net, red, use):
    data = load_targets()
    data[month_key] = {"net": int(net), "red": int(red), "use": int(use)}
    keys = sorted(data.keys())
    for k in keys[:-12]:
        del data[k]
    try:
        with open(TARGETS_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f)
        return True
    except Exception:
        return False


def report_file_name(now=None):
    # file name = brand + yesterday's date (day-month), Iraq time (UTC+3)
    # example: today is the 11th of October -> "brand 10-10.xlsx"
    if now is None:
        now = datetime.now(timezone(timedelta(hours=3)))
    day = now - timedelta(days=1)
    return T["brand"] + " " + str(day.day) + "-" + str(day.month) + ".xlsx"


def show_table(df):
    # shows the table on screen with the same look as the Excel file:
    # gray cells, yellow first column, borders, centered, equal widths, font 14
    if len(df) > 3000:
        st.dataframe(df, use_container_width=True)
        return
    try:
        def fmt(v):
            if isinstance(v, float) and v == v and v == int(v):
                return str(int(v))
            return v

        sty = style_screen(df)
        sty = sty.hide(axis="index")
        sty = sty.format(fmt)
        total_w = 160 * len(df.columns)
        sty = sty.set_table_styles([
            {"selector": "table", "props": [
                ("border-collapse", "collapse"),
                ("table-layout", "fixed"),
                ("width", str(total_w) + "px"),
            ]},
            {"selector": "th", "props": [
                ("background-color", "#A6A6A6"),
                ("color", "black"),
                ("border", "1px solid black"),
                ("text-align", "center"),
                ("font-size", "14px"),
                ("padding", "6px"),
                ("word-wrap", "break-word"),
            ]},
            {"selector": "th.col_heading.col0", "props": [("background-color", "#FFFF00")]},
            {"selector": "td", "props": [("padding", "6px"), ("word-wrap", "break-word")]},
        ])
        html = sty.to_html()
        st.markdown('<div style="overflow-x:auto">' + html + "</div>", unsafe_allow_html=True)
    except Exception:
        st.dataframe(df, use_container_width=True)


st.title(T["title"])

c1, c2, c3 = st.columns(3)
c4, c5, c6 = st.columns(3)
with c1:
    agents_file = st.file_uploader(T["up_agents"], key="agents")
with c2:
    dist_file = st.file_uploader(T["up_dist"], key="dist")
with c3:
    quality_file = st.file_uploader(T["up_quality"], key="quality")
with c4:
    inc_file = st.file_uploader(T["up_inc"], key="inc")
with c5:
    wal_file = st.file_uploader(T["up_wal"], key="wal")
with c6:
    bnd_file = st.file_uploader(T["up_bnd"], key="bnd")

mk = target_month_key()
saved_t = load_targets().get(mk, {})
with st.expander(T["tgt_title"] + " - " + mk, expanded=True):
    t1, t2, t3 = st.columns(3)
    with t1:
        tgt_net = st.number_input(T["tgt_net"], min_value=0, value=int(saved_t.get("net", 0)), step=100, key="tgt_net_" + mk)
    with t2:
        tgt_red = st.number_input(T["tgt_red"], min_value=0, value=int(saved_t.get("red", 0)), step=100, key="tgt_red_" + mk)
    with t3:
        tgt_use = st.number_input(T["tgt_use"], min_value=0, value=int(saved_t.get("use", 0)), step=100, key="tgt_use_" + mk)

if agents_file is not None and dist_file is not None:
    agents = None
    dist = None
    q = None
    inc = None
    wal = None
    bnd = None

    with c1:
        sh = load_report(agents_file, T["agents"])
        if sh is not None:
            agents = pick_sheet(sh, T["agents"])
    with c2:
        sh = load_report(dist_file, T["dist"])
        if sh is not None:
            dist = pick_sheet(sh, T["dist"])
    if quality_file is not None:
        with c3:
            sh = load_report(quality_file, T["quality"])
            if sh is not None:
                q = pick_sheet(sh, T["quality"])
    if inc_file is not None:
        with c4:
            sh = load_report(inc_file, T["inc"])
            if sh is not None:
                inc = pick_sheet(sh, T["inc"])
    if wal_file is not None:
        with c5:
            sh = load_report(wal_file, T["wal"])
            if sh is not None:
                wal = pick_sheet(sh, T["wal"])
    if bnd_file is not None:
        with c6:
            sh = load_report(bnd_file, T["bnd"])
            if sh is not None:
                bnd = pick_sheet(sh, T["bnd"])

    if agents is not None and dist is not None:
        with st.expander(T["preview"]):
            st.write(T["agents"])
            st.dataframe(agents.head(5))
            st.write(T["dist"])
            st.dataframe(dist.head(5))
            if q is not None:
                st.write(T["quality"])
                st.dataframe(q.head(5))
            if inc is not None:
                st.write(T["inc"])
                st.dataframe(inc.head(5))
            if wal is not None:
                st.write(T["wal"])
                st.dataframe(wal.head(5))
            if bnd is not None:
                st.write(T["bnd"])
                st.dataframe(bnd.head(5))

        if st.button(T["calc"]):
            save_ok = save_targets(mk, tgt_net, tgt_red, tgt_use)
            if not save_ok:
                st.caption(T["tgt_nosave"])
            try:
                result = compute(agents, dist, q, inc, wal, bnd)
                if result is not None:
                    show_table(result)
                    bundle_df = result.attrs.get("bundles")
                    mm_df = result.attrs.get("mismatch")
                    extras = []
                    after_points = []
                    if bundle_df is not None:
                        st.write(T["bnd"])
                        if len(bundle_df) > 0:
                            show_table(bundle_df)
                        else:
                            st.info(T["bnd_empty"])
                        extras.append((T["bnd"], bundle_df))
                    if mm_df is not None and len(mm_df) > 0:
                        st.write(T["mm_title"])
                        show_table(mm_df)
                        extras.append((T["mm_sheet"], mm_df))
                    monthly_df = result.attrs.get("monthly")
                    if monthly_df is not None:
                        m_name = T["monthly_sheet"] + " " + str(result.attrs.get("monthly_month"))
                        st.write(m_name)
                        show_table(monthly_df)
                        after_points = [(m_name, monthly_df)]
                        if result.attrs.get("monthly_msg"):
                            st.info(result.attrs.get("monthly_msg"))
                    points_info = None
                    points_src = result.attrs.get("points_src")
                    if points_src is not None:
                        points_info = build_points_table(points_src, tgt_net, tgt_red, tgt_use)
                    if points_info is not None:
                        st.write(T["pts_sheet"])
                        st.dataframe(points_info["display"], use_container_width=True)
                    xlsx_bytes = format_excel(result, extras, points_info, after_points)
                    st.download_button(
                        T["download"],
                        data=xlsx_bytes,
                        file_name=report_file_name(),
                        mime=XLSX_MIME,
                    )
            except Exception:
                st.error(T["err"])
                st.code(traceback.format_exc())
